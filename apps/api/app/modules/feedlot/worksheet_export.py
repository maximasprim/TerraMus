"""Builds the Sahara Oasis "Bulls WG Worksheet" from live data, in the worksheet's exact layout.

Same sheet names/order, same column order, same headers, yellow header/TOTAL styling,
thin borders, and the same colour rules - but with correct, live Excel formulas:
  * dates are real dates shown as dd/mm/yyyy (the original's m/d/yy caused swaps)
  * day counts / bull counts are derived, never hand-typed
  * total-row % gain divides by the entry weight of the SAME animals
  * the profit table counts every sold bull, charges culls as losses, and allocates
    operations cost per bull-day instead of a flat figure
Added value goes in extra columns to the right and in extra sheets AFTER the original ones.
Importing the exported file again reproduces the same data (round trip).
"""
from __future__ import annotations

import re
import uuid
from collections import defaultdict
from datetime import date
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.modules.feedlot import metrics, service
from app.modules.feedlot.models import IntakeBatch, SaleBatch

YELLOW = PatternFill("solid", fgColor="FFFFFF00")
GREEN, CYAN, RED = (PatternFill("solid", bgColor=c, fgColor=c) for c in ("FF00FF00", "FF00FFFF", "FFFF0000"))
THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
F = Font(name="Arial", size=11)
FB = Font(name="Arial", size=11, bold=True)
DATE_FMT = "dd/mm/yyyy"
SALES_NAMES = {1: " Sales weight Batch 1", 2: " Sales weight Batch 2", 3: "Sales weight Batch 3",
               4: "Sales weight Batch4", 5: "Sales Weight Batch 5", 6: "Sales  weight Batch 6 "}
FEED_ROWS = (("silage", "SILAGE"), ("hay", "HAY"), ("concentrate", "CONCENTRATE"))


def _nat(tag: str):
    return [int(p) if p.isdigit() else p.lower() for p in re.split(r"(\d+)", tag)]


def intake_sheet_name(n: int) -> str:
    return "BATCH 4." if n == 4 else f"BATCH {n}"


def sales_sheet_name(n: int) -> str:
    return SALES_NAMES.get(n, f"Sales weight Batch {n}")


def q(name: str) -> str:
    return "'" + name.replace("'", "''") + "'"


def put(ws, r, c, v=None, *, nf=None, bold=False, fill=None, wrap=False, border=True, align=None):
    cell = ws.cell(r, c, v)
    cell.font = FB if bold else F
    if border:
        cell.border = BORDER
    if fill:
        cell.fill = fill
    if nf:
        cell.number_format = nf
    if wrap or align:
        cell.alignment = Alignment(wrap_text=wrap, horizontal=align, vertical="bottom")
    return cell


def hdr(ws, r, c, v, **kw):
    return put(ws, r, c, v, fill=YELLOW, wrap=True, **kw)


def dexpr(d: date) -> str:
    return f"DATE({d.year},{d.month},{d.day})"


def _color_rules(ws, rng: str, first_cell: str, kind: str, s: dict) -> None:
    n = f"ISNUMBER({first_cell})"
    if kind == "weight":
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f"AND({n},{first_cell}>={s['market_ready_kg']})"], fill=GREEN))
    elif kind == "gain":
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f"AND({n},{first_cell}>{s['interval_gain_good']})"], fill=CYAN))
    elif kind == "adg":
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f"AND({n},{first_cell}<={s['adg_poor_max']})"], fill=RED))
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f"AND({n},{first_cell}>={s['adg_target_min']},{first_cell}<={s['adg_target_max']})"], fill=GREEN))
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f"AND({n},{first_cell}>{s['adg_target_max']})"], fill=CYAN))


# ----------------------------------------------------------------- data bundle
class Bundle:
    def __init__(self):
        self.as_of: date = date.today()
        self.s: dict = metrics.get_settings(None)
        self.batches: list[IntakeBatch] = []
        self.sale_batches: list[SaleBatch] = []
        self.items: list[tuple[metrics.AnimalInput, dict]] = []
        self.docs: dict[str, list[Any]] = defaultdict(list)
        self.profit: dict = {}
        self.perf: dict = {}


async def gather(db: AsyncSession, org_id: uuid.UUID, as_of: date | None) -> Bundle:
    b = Bundle()
    b.s = metrics.get_settings(await service._settings_data(db, org_id))
    b.batches = list((await db.execute(select(IntakeBatch).where(IntakeBatch.organization_id == org_id).order_by(IntakeBatch.batch_number))).scalars())
    b.sale_batches = list((await db.execute(select(SaleBatch).where(SaleBatch.organization_id == org_id).order_by(SaleBatch.batch_number))).scalars())
    b.items = await service._load(db, org_id)
    for d in await service.list_costing_docs(db, org_id):
        b.docs[d.kind].append(d)
    all_w = [d for inp, _ in b.items for d, _ in inp.weights]
    ex_d = [inp.exit_date for inp, _ in b.items if inp.exit_date]
    b.as_of = as_of or max(all_w + ex_d + [date.today()] if not (all_w or ex_d) else all_w + ex_d)
    b.profit = await service.profitability(db, org_id, b.as_of)
    b.perf = await service.performance(db, org_id, b.as_of)
    return b


# ------------------------------------------------------------------ BATCH n
def build_batch_sheet(wb, b: Bundle, batch: IntakeBatch, animals: list[metrics.AnimalInput]) -> dict:
    ws = wb.create_sheet(intake_sheet_name(batch.batch_number))
    off = int(b.s["day_offset"])
    W = sorted({d for a in animals for d, _ in a.weights if d <= b.as_of})
    ref = {"sheet": ws.title, "first_row": 3, "last_row": 2 + len(animals), "total_row": 3 + len(animals)}
    wcols, gcols, c = {}, {}, 2
    if W:
        wcols[W[0]] = 2
        c = 3
        for k in range(1, len(W)):
            wcols[W[k]], gcols[W[k]] = c, c + 1
            c += 2
    have_overall = len(W) >= 2
    ov = dw = pc = None
    if have_overall:
        ov, dw, pc = c, c + 1, c + 2
    last_col = (pc or (c - 1 if W else 2))
    title = f"BATCH {batch.batch_number}."
    if batch.source:
        title += f" {batch.source.upper()}- {batch.head_purchased or len(animals)} {batch.sex_class.title()}s."
    title += f" Delivered on {batch.delivered_on.strftime('%d/%m/%Y')}."
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(last_col, 2))
    put(ws, 1, 1, title, bold=True, align="center")
    hdr(ws, 2, 1, "ANIMAL NO.")
    for d, col in wcols.items():
        put(ws, 2, col, d, nf=DATE_FMT, fill=YELLOW, wrap=True)
    prev = W[0] if W else None
    for k in range(1, len(W)):
        hdr(ws, 2, gcols[W[k]], f"{(W[k] - W[k - 1]).days} Day Gain/ (Loss)")
    asof_lbl = W[-1].strftime("%d/%m/%Y") if W else ""
    days_expr = f"{dexpr(W[-1])}-{dexpr(batch.delivered_on)}-{off}" if W else "1"
    if have_overall:
        hdr(ws, 2, ov, f"Overall Gain/loss {asof_lbl}"); hdr(ws, 2, dw, f"Daily WG {asof_lbl}"); hdr(ws, 2, pc, f"% WG as at {asof_lbl}")
    ws.row_dimensions[2].height = 45
    for i, a in enumerate(sorted(animals, key=lambda x: _nat(x.tag))):
        r = 3 + i
        put(ws, r, 1, a.tag)
        wmap = {d: kg for d, kg in a.weights}
        for d, col in wcols.items():
            put(ws, r, col, wmap.get(d), align="right")
        for k in range(1, len(W)):
            cur, prv = L(wcols[W[k]]), L(wcols[W[k - 1]])
            put(ws, r, gcols[W[k]], f'=IF(AND(ISNUMBER({cur}{r}),ISNUMBER({prv}{r})),{cur}{r}-{prv}{r},"")', align="right")
        if have_overall:
            lw = L(wcols[W[-1]])
            put(ws, r, ov, f'=IF(AND(ISNUMBER({lw}{r}),ISNUMBER(B{r})),{lw}{r}-B{r},"")')
            put(ws, r, dw, f'=IF(ISNUMBER({L(ov)}{r}),{L(ov)}{r}/({days_expr}),"")', nf="0.00")
            put(ws, r, pc, f'=IF(ISNUMBER({L(ov)}{r}),{L(ov)}{r}/B{r},"")', nf="0.0%")
    t, f_, l_ = ref["total_row"], 3, ref["last_row"]
    put(ws, t, 1, "TOTAL", bold=True, fill=YELLOW)
    for col in list(wcols.values()) + list(gcols.values()) + ([ov] if ov else []):
        put(ws, t, col, f"=SUM({L(col)}{f_}:{L(col)}{l_})", bold=True, fill=YELLOW)
    if have_overall:
        rng = f"{L(ov)}{f_}:{L(ov)}{l_}"
        put(ws, t, dw, f'=IFERROR(SUM({rng})/({days_expr})/COUNT({rng}),"")', bold=True, fill=YELLOW, nf="0.00")
        put(ws, t, pc, f'=IFERROR(SUM({rng})/SUMPRODUCT(--ISNUMBER({rng}),B{f_}:B{l_}),"")', bold=True, fill=YELLOW, nf="0.0%")
        ref.update(overall_col=L(ov), adg_cell=f"{L(dw)}{t}", pct_cell=f"{L(pc)}{t}")
        _color_rules(ws, f"{L(dw)}{f_}:{L(dw)}{l_}", f"{L(dw)}{f_}", "adg", b.s)
    for col in wcols.values():
        _color_rules(ws, f"{L(col)}{f_}:{L(col)}{l_}", f"{L(col)}{f_}", "weight", b.s)
    for col in gcols.values():
        _color_rules(ws, f"{L(col)}{f_}:{L(col)}{l_}", f"{L(col)}{f_}", "gain", b.s)
    ws.column_dimensions["A"].width = 12
    for col in range(2, last_col + 1):
        ws.column_dimensions[L(col)].width = 12
    ws.freeze_panes = "B3"
    ref["as_of"] = W[-1] if W else None
    return ref


# --------------------------------------------------------------- sales sheets
def build_sales_sheet(wb, b: Bundle, sb: SaleBatch, rows: list[metrics.AnimalInput], batch_no: dict) -> dict:
    ws = wb.create_sheet(sales_sheet_name(sb.batch_number))
    heads = ["Animal No.", "Date received", "Entry weight(Kg)", "Date sold", "Sales Weight(Kg)", "TWG", "Days on Farm", "ADG",
             "Last on-farm weight(Kg)", "Last weigh date", "Sale vs last weigh (Kg)", "Price per Kg (KES)", "Sale amount (KES)"]
    for i, h in enumerate(heads, 1):
        put(ws, 1, i, h, bold=True, fill=YELLOW, wrap=True)
    r = 2
    last_group = None
    groups = sorted(rows, key=lambda a: (a.intake_batch or 0, _nat(a.tag)))
    off = int(b.s["day_offset"])
    for a in groups:
        if last_group is not None and a.intake_batch != last_group:
            put(ws, r, 4, f"BATCH {a.intake_batch}", bold=True, border=False)
            r += 1
        last_group = a.intake_batch
        prior = sorted((d, kg) for d, kg in a.weights if d <= a.exit_date)
        put(ws, r, 1, a.tag, bold=True)
        put(ws, r, 2, a.received_on, nf=DATE_FMT, bold=True)
        put(ws, r, 3, a.entry_weight)
        put(ws, r, 4, a.exit_date, nf=DATE_FMT)
        put(ws, r, 5, a.exit_weight)
        put(ws, r, 6, f"=E{r}-C{r}")
        put(ws, r, 7, f"=D{r}-B{r}-{off}")
        put(ws, r, 8, f'=IF(G{r}>0,F{r}/G{r},"")', nf="0.00")
        put(ws, r, 9, prior[-1][1] if prior else None)
        put(ws, r, 10, prior[-1][0] if prior else None, nf=DATE_FMT)
        put(ws, r, 11, f'=IF(ISNUMBER(I{r}),E{r}-I{r},"")')
        put(ws, r, 12, float(sb.price_per_kg) if sb.price_per_kg else (a.exit_amount / a.exit_weight if a.exit_weight else None), nf="0.00")
        put(ws, r, 13, f"=E{r}*L{r}", nf="#,##0.00")
        r += 1
    first, last = 2, r - 1
    put(ws, r, 1, f"=COUNTA(A{first}:A{last})", bold=True, fill=YELLOW)
    put(ws, r, 5, f"=SUM(E{first}:E{last})", bold=True, fill=YELLOW)
    put(ws, r, 13, f"=SUM(M{first}:M{last})", bold=True, fill=YELLOW, nf="#,##0.00")
    _color_rules(ws, f"H{first}:H{last}", f"H{first}", "adg", b.s)
    for c_, w in zip("ABCDEFGHIJKLM", (12, 13, 14, 12, 15, 8, 11, 8, 15, 13, 14, 13, 16)):
        ws.column_dimensions[c_].width = w
    ws.freeze_panes = "A2"
    return {"sheet": ws.title, "first": first, "last": last, "total_row": r}


# ----------------------------------------------------------------- the register
def build_register(wb, b: Bundle, by_batch: dict[int, list[metrics.AnimalInput]]) -> None:
    ws = wb.create_sheet("Bull Inventory register")
    nums = sorted(by_batch)
    if not nums:
        put(ws, 1, 2, "SOLD BULLS", bold=True, border=False)
        put(ws, 1, 8, "CURRENT STOCK", bold=True, border=False)
        return
    for i, n in enumerate(nums):
        hdr(ws, 2, 1 + i, f"Batch {n}")
    put(ws, 1, 2, "SOLD BULLS", bold=True, border=False)
    sold_cols = {}
    maxr = 2
    for i, n in enumerate(nums):
        r = 3
        for a in sorted(by_batch[n], key=lambda x: _nat(x.tag)):
            if a.exit_type == "sold":
                put(ws, r, 1 + i, f"{a.tag}...{a.sale_batch}"); r += 1
            elif a.exit_type == "culled":
                put(ws, r, 1 + i, f"{a.tag}...culled"); r += 1
            elif a.exit_type == "died":
                put(ws, r, 1 + i, f"{a.tag}...died"); r += 1
        sold_cols[n] = (1 + i, r - 1)
        maxr = max(maxr, r - 1)
    stock_start = len(nums) + 2
    stock_batches = [n for n in nums if any(a.exit_type is None for a in by_batch[n])]
    if stock_batches:
        put(ws, 1, stock_start + len(stock_batches) // 2, "CURRENT STOCK", bold=True, border=False)
    stock_total = 0
    for j, n in enumerate(stock_batches):
        hdr(ws, 2, stock_start + j, f"Batch {n}")
        r = 3
        for a in sorted((x for x in by_batch[n] if x.exit_type is None), key=lambda x: _nat(x.tag)):
            put(ws, r, stock_start + j, a.tag); r += 1
        stock_total += r - 3
        maxr = max(maxr, r - 1)
    t = maxr + 1
    for i, n in enumerate(nums):
        put(ws, t, 1 + i, f"=COUNTA({L(1 + i)}3:{L(1 + i)}{maxr})", bold=True, fill=YELLOW)
    put(ws, t, len(nums) + 1, f"=SUM(A{t}:{L(len(nums))}{t})", bold=True, fill=YELLOW)
    for j in range(len(stock_batches)):
        put(ws, t, stock_start + j, f"=COUNTA({L(stock_start + j)}3:{L(stock_start + j)}{maxr})", bold=True, fill=YELLOW)
    if stock_batches:
        put(ws, t, stock_start + len(stock_batches), f"=SUM({L(stock_start)}{t}:{L(stock_start + len(stock_batches) - 1)}{t})", bold=True, fill=YELLOW)
    for col in range(1, stock_start + len(stock_batches) + 2):
        ws.column_dimensions[L(col)].width = 13


# ---------------------------------------------------------------- cost sheets
def _doc(b: Bundle, kind: str, idx: int = 0):
    docs = b.docs.get(kind, [])
    return docs[idx] if len(docs) > idx else None


def build_concentrate(wb, b: Bundle) -> dict:
    ws = wb.create_sheet("CONCENTRATE")
    doc = _doc(b, "recipe")
    data = doc.data if doc else {"ingredients": [], "fixed_costs": []}
    put(ws, 1, 1, "CONCENTRATE ONE TONNE COMPOSITION FEEDS.", bold=True, border=False)
    for i, h in enumerate(["COST ITEM", "QUANTITY(Kg)", "Price per Kg", "Total Cost."], 1):
        hdr(ws, 2, i, h)
    r = 3
    for ing in data.get("ingredients", []):
        put(ws, r, 1, ing["name"]); put(ws, r, 2, ing["quantity_kg"]); put(ws, r, 3, ing["price_per_kg"]); put(ws, r, 4, f"=B{r}*C{r}")
        r += 1
    qty_last = r - 1
    for fc in data.get("fixed_costs", []):
        put(ws, r, 1, fc["name"]); put(ws, r, 4, fc["amount"]); r += 1
    put(ws, r, 1, "TOTAL", bold=True, fill=YELLOW)
    put(ws, r, 2, f"=SUM(B3:B{qty_last})" if qty_last >= 3 else 0, bold=True, fill=YELLOW)
    put(ws, r, 4, f"=SUM(D3:D{r - 1})" if r > 3 else 0, bold=True, fill=YELLOW)
    put(ws, r, 3, f'=IFERROR(D{r}/B{r},0)', bold=True, fill=YELLOW, nf="0.00")
    put(ws, r + 2, 3, f'=IFERROR(D{r}/B{r},0)', nf="0.00", bold=True)
    ws.column_dimensions["A"].width = 30
    for c_ in "BCD":
        ws.column_dimensions[c_].width = 15
    return {"cost_cell": f"CONCENTRATE!C{r + 2}"}


def _title_for(name: str) -> str:
    return re.sub(r"\s*\(per acre\)\s*", "", name, flags=re.I).upper() + " (PER ACRE)"


def _budget_block(ws, top: int, name: str, data: dict, unit_label: str) -> dict:
    put(ws, top, 2, _title_for(name), bold=True, border=False)
    put(ws, top + 1, 2, "Assumptions", bold=True, border=False)
    for i, a in enumerate((data.get("assumptions") or [])[:6]):
        put(ws, top + 2 + i, 2, a, border=False)
    h = top + 8
    hdr(ws, h, 2, "Cost Item"); hdr(ws, h, 3, "Total Cost"); hdr(ws, h, 4, "Per Year")
    r = h + 1
    for i, ln in enumerate(data.get("lines", []), 1):
        put(ws, r, 1, i); put(ws, r, 2, ln["item"]); put(ws, r, 3, ln["amount"])
        mult, div = float(ln.get("multiply_by", 1) or 1), float(ln.get("divide_by", 1) or 1)
        fx = f"=C{r}/{div:g}" if div != 1 else f"=C{r}*{mult:g}" if mult != 1 else f"=C{r}"
        put(ws, r, 4, fx)
        r += 1
    put(ws, r, 4, f"=SUM(D{h + 1}:D{r - 1})", bold=True)
    put(ws, r + 1, 2, f"Total Yield per year ({unit_label})", bold=True); put(ws, r + 1, 4, data.get("yield_per_year"))
    put(ws, r + 2, 2, "Cost per bale" if data.get("unit_weight_kg") else "Cost per Kg", bold=True)
    put(ws, r + 2, 4, f"=IFERROR(D{r}/D{r + 1},0)", nf="0.00")
    out = {"per_year_total": f"D{r}", "yield": f"D{r + 1}", "lines": (h + 1, r - 1), "end": r + 2}
    if data.get("unit_weight_kg"):
        put(ws, r + 3, 2, "Cost per kg", bold=True)
        put(ws, r + 3, 4, f"=D{r + 2}/{data['unit_weight_kg']:g}", nf="0.00")
        out["per_kg_cell"] = f"D{r + 3}"; out["end"] = r + 3
    else:
        out["per_kg_cell"] = f"D{r + 2}"
    return out


def build_hay(wb, b: Bundle) -> dict:
    ws = wb.create_sheet("HAY")
    name = next((d.name for d in b.docs.get("enterprise_budget", []) if "hay" in d.name.lower()), None)
    doc = next((d for d in b.docs.get("enterprise_budget", []) if d.name == name), None)
    out = {"cost_cell": "HAY!D24"}
    if doc:
        blk = _budget_block(ws, 1, doc.name, doc.data, "bales" if doc.data.get("unit_weight_kg") else "kg")
        out["cost_cell"] = f"HAY!{blk['per_kg_cell']}"
    ws.column_dimensions["A"].width = 5; ws.column_dimensions["B"].width = 46
    for c_ in "CD":
        ws.column_dimensions[c_].width = 14
    return out


def build_silage(wb, b: Bundle) -> dict:
    ws = wb.create_sheet("SILAGE -BLENDED")
    docs = [d for d in b.docs.get("enterprise_budget", []) if "hay" not in d.name.lower()]
    top, blocks = 2, []
    for d in docs:
        blk = _budget_block(ws, top, d.name, d.data, "kgs")
        blocks.append((d, blk))
        top = blk["end"] + 4
    out = {"cost_cell": "'SILAGE -BLENDED'!F98"}
    if not blocks:
        return out
    r0 = top
    put(ws, r0, 2, "SILAGE- BLENDED", bold=True, border=False)
    hdr(ws, r0 + 1, 2, "Cost Item")
    for j, (d, _) in enumerate(blocks):
        hdr(ws, r0 + 1, 3 + j, re.sub(r"\s*\(per acre\)", "", d.name, flags=re.I))
    ftot = 3 + len(blocks)
    hdr(ws, r0 + 1, ftot, "Total cost (per year)")
    items = []
    for d, _ in blocks:
        for ln in d.data.get("lines", []):
            if ln["item"] not in items:
                items.append(ln["item"])
    r = r0 + 2
    for i, it in enumerate(items, 1):
        put(ws, r, 1, i); put(ws, r, 2, it)
        for j, (d, blk) in enumerate(blocks):
            idx = next((k for k, ln in enumerate(d.data.get("lines", [])) if ln["item"] == it), None)
            put(ws, r, 3 + j, f"=D{blk['lines'][0] + idx}" if idx is not None else 0)
        put(ws, r, ftot, f"=SUM(C{r}:{L(ftot - 1)}{r})")
        r += 1
    cpa, tc, ypa, ty, cpk = r + 1, r + 2, r + 4, r + 5, r + 6
    put(ws, cpa, 2, "Cost per acre per year", bold=True); put(ws, tc, 2, "Total Cost ", bold=True)
    put(ws, ypa, 2, "Yield per acre per year (kgs)", bold=True); put(ws, ty, 2, "Total Yield per year (kgs)", bold=True)
    put(ws, cpk, 2, "Cost per Kg", bold=True)
    for j, (d, blk) in enumerate(blocks):
        col = L(3 + j)
        acres = d.data.get("acres") or 1
        put(ws, cpa, 3 + j, f"=SUM({col}{r0 + 2}:{col}{r - 1})")
        put(ws, tc, 3 + j, f"={col}{cpa}*{acres:g}")
        put(ws, ypa, 3 + j, d.data.get("yield_per_year"))
        put(ws, ty, 3 + j, f"={col}{ypa}*{acres:g}")
        put(ws, cpk, 3 + j, f"=IFERROR({col}{tc}/{col}{ty},0)", nf="0.00")
    fcol = L(ftot); lastc = L(ftot - 1)
    put(ws, cpa, ftot, f"=SUM({fcol}{r0 + 2}:{fcol}{r - 1})")
    put(ws, tc, ftot, f"=SUM(C{tc}:{lastc}{tc})"); put(ws, ypa, ftot, f"=SUM(C{ypa}:{lastc}{ypa})")
    put(ws, ty, ftot, f"=SUM(C{ty}:{lastc}{ty})"); put(ws, cpk, ftot, f"=IFERROR({fcol}{tc}/{fcol}{ty},0)", nf="0.00", bold=True)
    ws.column_dimensions["A"].width = 5; ws.column_dimensions["B"].width = 44
    for c_ in range(3, ftot + 1):
        ws.column_dimensions[L(c_)].width = 15
    out["cost_cell"] = f"'SILAGE -BLENDED'!{fcol}{cpk}"
    return out


def build_land(wb, b: Bundle) -> None:
    ws = wb.create_sheet("LAND USE")
    doc = _doc(b, "land_use")
    put(ws, 1, 1, "LAND USE IN ACRES", bold=True, border=False)
    sites = (doc.data.get("sites") if doc else None) or []
    hdr(ws, 2, 1, "ITEM")
    for j, sname in enumerate(sites):
        hdr(ws, 2, 2 + j, sname)
    tcol = 2 + len(sites)
    hdr(ws, 2, tcol, "TOTAL")
    r = 3
    for row in (doc.data.get("rows", []) if doc else []):
        put(ws, r, 1, row["item"])
        for j, sname in enumerate(sites):
            put(ws, r, 2 + j, row.get(sname, 0))
        put(ws, r, tcol, f"=SUM(B{r}:{L(tcol - 1)}{r})")
        put(ws, r, tcol + 1, "Irrigated" if row.get("irrigated") else "Not Irrigated ", border=False)
        r += 1
    put(ws, r, 1, "TOTAL", bold=True, fill=YELLOW)
    for j in range(len(sites) + 1):
        put(ws, r, 2 + j, f"=SUM({L(2 + j)}3:{L(2 + j)}{r - 1})", bold=True, fill=YELLOW)
    ws.column_dimensions["A"].width = 20


def build_ration(wb, b: Bundle, refs: dict) -> dict:
    ws = wb.create_sheet("TOTAL FEEDING COSTS")
    doc = _doc(b, "ration_plan")
    stages = doc.data.get("stages", []) if doc else []
    total_days = sum(int(s.get("days", 0)) for s in stages)
    put(ws, 2, 1, f"DAILY CONSUMPTION ON {total_days} DAY FATTENING", bold=True, border=False)
    for j, st in enumerate(stages):
        hdr(ws, 4, 2 + j, st.get("label", f"Stage {j + 1}"))
    row_of = {}
    for i, (key, label) in enumerate(FEED_ROWS):
        row_of[key] = 5 + i
        put(ws, 5 + i, 1, label, bold=True)
        for j, st in enumerate(stages):
            put(ws, 5 + i, 2 + j, (st.get("feeds") or {}).get(key, 0))
    put(ws, 10, 2, f"FEED COST  PER BULL FOR A {total_days}-DAY FATTENING PERIOD", bold=True, border=False)
    for j, h in enumerate(["Period ", "Hay (kg/day)", "Silage (kg/day)", "Concentrate (kg/day)", "Hay Cost/Day", "Silage Cost/Day",
                           "Concentrate Cost/Day", "Total Cost/Day", "Days", "Start day", "Cumulative cost"], 1):
        hdr(ws, 11, j, h)
    r0 = 12
    for j, st in enumerate(stages):
        r, sc = r0 + j, L(2 + j)
        put(ws, r, 1, st.get("label", f"Stage {j + 1}"))
        put(ws, r, 2, f"={sc}6"); put(ws, r, 3, f"={sc}5"); put(ws, r, 4, f"={sc}7")
        put(ws, r, 5, f"=B{r}*{refs['hay']}"); put(ws, r, 6, f"=C{r}*{refs['silage']}"); put(ws, r, 7, f"=D{r}*{refs['concentrate']}")
        put(ws, r, 8, f"=SUM(E{r}:G{r})"); put(ws, r, 9, int(st.get("days", 0)))
        put(ws, r, 10, 0 if j == 0 else f"=J{r - 1}+I{r - 1}")
        put(ws, r, 11, f"=H{r}*I{r}" if j == 0 else f"=K{r - 1}+H{r}*I{r}")
    rl = r0 + max(len(stages), 1) - 1
    put(ws, rl + 1, 1, "Total", bold=True, fill=YELLOW)
    put(ws, rl + 1, 9, f"=SUM(I{r0}:I{rl})", bold=True, fill=YELLOW)
    put(ws, rl + 1, 11, f"=K{rl}" if stages else 0, bold=True, fill=YELLOW)
    ws.column_dimensions["A"].width = 24
    for c_ in range(2, 12):
        ws.column_dimensions[L(c_)].width = 15
    q_ = "'TOTAL FEEDING COSTS'!"
    return {"has_plan": bool(stages), "H": f"{q_}$H${r0}:$H${rl}", "I": f"{q_}$I${r0}:$I${rl}", "J": f"{q_}$J${r0}:$J${rl}",
            "total_days": f"{q_}$I${rl + 1}", "last_cost": f"{q_}$H${rl}"}


def feed_formula(rt: dict, d: str) -> str:
    if not rt["has_plan"]:
        return "0"
    x = f"({d}-{rt['J']})"
    m = f"((ABS{x}+{x})/2)"
    return f"SUMPRODUCT(({m}+{rt['I']}-ABS({m}-{rt['I']}))/2,{rt['H']})+MAX({d}-{rt['total_days']},0)*{rt['last_cost']}"


def build_ops(wb, b: Bundle, bull_days: int) -> dict:
    ws = wb.create_sheet("OPERATIONS COST")
    doc = _doc(b, "overhead_pool")
    lines = doc.data.get("lines", []) if doc else []
    heads = int(doc.data.get("heads") or 0) if doc else 0
    cycle = int(doc.data.get("cycle_days") or 0) if doc else 0
    put(ws, 1, 1, "BULL FATTENING  OPERATIONAL COSTS (EXCLUDING  FEED PRODUCTION)", bold=True, border=False)
    hdr(ws, 2, 1, "Cost Category"); hdr(ws, 2, 2, "Total Cost")
    r = 3
    for ln in lines:
        put(ws, r, 1, ln["item"]); put(ws, r, 2, ln["amount"]); r += 1
    put(ws, r, 1, "Total Operational Cost", bold=True, fill=YELLOW); put(ws, r, 2, f"=SUM(B3:B{r - 1})" if lines else 0, bold=True, fill=YELLOW)
    top_total = r
    put(ws, r + 1, 2, f"=IFERROR(B{r}/{heads},0)" if heads else 0, nf="#,##0.00")
    t0 = r + 3
    put(ws, t0, 1, "BULL FATTENING  OPERATIONAL COSTS (EXCLUDING  FEED PRODUCTION)- FOR ONE PRODUCTION CYCLE.", bold=True, border=False)
    hdr(ws, t0 + 1, 1, "Cost Category"); hdr(ws, t0 + 1, 2, "Cost Behaviour"); hdr(ws, t0 + 1, 3, f"Historical Cost ({heads} Bulls)")
    rr = t0 + 2
    order = [("fixed", "Fixed Costs"), ("variable", "Variable Costs"), ("step", "Semi-variable Costs"), ("semi-variable", "Semi-variable Costs")]
    done = set()
    sub_rows = []
    for beh, label in order:
        if label in done:
            continue
        grp = [ln for ln in lines if (ln.get("behaviour") or "fixed") in ([beh] if label != "Semi-variable Costs" else ["step", "semi-variable"])]
        if not grp:
            continue
        done.add(label)
        put(ws, rr, 1, label, bold=True, border=False); rr += 1
        first = rr
        for ln in grp:
            src = 3 + lines.index(ln)
            put(ws, rr, 1, ln["item"])
            put(ws, rr, 2, {"fixed": "Fixed", "variable": "Variable", "step": "Step Cost", "semi-variable": "Semi-variable"}.get(ln.get("behaviour"), "Fixed"))
            put(ws, rr, 3, f"=B{src}")
            rr += 1
        put(ws, rr, 1, f"Subtotal – {label}", bold=True); put(ws, rr, 3, f"=SUM(C{first}:C{rr - 1})", bold=True)
        sub_rows.append(rr); rr += 2
    put(ws, rr, 1, "Total Operational Cost ", bold=True, fill=YELLOW)
    put(ws, rr, 3, "=" + "+".join(f"C{x}" for x in sub_rows) if sub_rows else 0, bold=True, fill=YELLOW)
    # cost summary block (F:G) - the cells the original used
    put(ws, t0 + 1, 6, "COST SUMMARY ", bold=True, border=False)
    hdr(ws, t0 + 2, 6, "Indicator"); hdr(ws, t0 + 2, 7, "Value")
    rows = [("Number of Bulls", heads), ("Production Cycle", cycle), ("Total Operational Cost", f"=C{rr}"),
            ("Operational Cost per Bull", f"=IFERROR(G{t0 + 5}/G{t0 + 3},0)"),
            ("Operational Cost per Bull per Day (planned cycle)", f"=IFERROR(G{t0 + 6}/G{t0 + 4},0)"),
            ("Actual bull-days to date", bull_days),
            ("Operational Cost per Bull per Day (actual bull-days)", f"=IFERROR(G{t0 + 5}/G{t0 + 8},0)")]
    for i, (k, v) in enumerate(rows):
        put(ws, t0 + 3 + i, 6, k); put(ws, t0 + 3 + i, 7, v, nf="#,##0.00" if i >= 3 else None)
    put(ws, t0 + 4, 8, "Days", border=False)
    ws.column_dimensions["A"].width = 52; ws.column_dimensions["B"].width = 16; ws.column_dimensions["C"].width = 24
    ws.column_dimensions["F"].width = 50; ws.column_dimensions["G"].width = 16
    return {"rate_cell": f"'OPERATIONS COST'!$G${t0 + 9}"}


# ------------------------------------------------- profitability (two sheets)
def build_profitability(wb, b: Bundle, by_batch, by_sale, refs: dict) -> None:
    prof = {a["animal_id"]: a for a in b.profit["animals"]}
    ws_pb = wb.create_sheet("Profitability per Bull")
    ws = wb.create_sheet(" PROFITABILITY ")
    # --- buying price table
    put(ws, 1, 4, "BUYING PRICE", bold=True, border=False)
    for i, h in enumerate(["Batch", "No of Bulls", "Cost", "Transport", "Other fees", "Handlers", "Total", "B.Price per Bull"], 1):
        hdr(ws, 2, i, h)
    buy_row = {}
    for i, bt in enumerate(b.batches):
        r = 3 + i
        buy_row[bt.batch_number] = r
        put(ws, r, 1, f"BATCH {bt.batch_number}"); put(ws, r, 2, bt.head_purchased)
        for c_, v in zip((3, 4, 5, 6), (bt.purchase_cost, bt.transport_cost, bt.other_fees, bt.handlers_cost)):
            put(ws, r, c_, float(v or 0))
        put(ws, r, 7, f"=SUM(C{r}:F{r})", nf="#,##0.00"); put(ws, r, 8, f"=IFERROR(G{r}/B{r},0)", nf="#,##0.00")
    s_top = 3 + len(b.batches) + 1
    put(ws, s_top, 4, "PROFITABILITY", bold=True, border=False)
    for i, h in enumerate(["", "No. of bulls", "Sales Weight(Kg)", "Selling price", "Buying price", "Feeding cost", "Operations cost", "Total expenses", "Profit"], 1):
        if h:
            hdr(ws, s_top + 1, i, h)
    # --- per bull blocks
    r = 1
    summary = []
    for sb in b.sale_batches:
        rows = by_sale.get(sb.batch_number)
        if not rows:
            continue
        put(ws_pb, r, 4, f"Sales Batch {sb.batch_number}", bold=True, border=False)
        for i, h in enumerate(["Animal No.", "Sales Weight(Kg)", "Days on Farm", "Selling price", "Buying price", "Feeding cost",
                               "Operations cost", "Total expenses", "Profit"], 1):
            hdr(ws_pb, r + 1, i, h)
        f = r + 2
        rate = float(sb.price_per_kg) if sb.price_per_kg else 0
        for a in sorted(rows, key=lambda x: (x.intake_batch or 0, _nat(x.tag))):
            p = prof.get(a.animal_id, {})
            put(ws_pb, f, 1, a.tag); put(ws_pb, f, 2, a.exit_weight); put(ws_pb, f, 3, p.get("days_on_farm", 0))
            put(ws_pb, f, 4, f"=B{f}*{rate:g}" if sb.price_per_kg else f"=B{f}*0", nf="#,##0.00")
            put(ws_pb, f, 5, f"={q(' PROFITABILITY ')}!H{buy_row.get(a.intake_batch, 3)}", nf="#,##0.00")
            put(ws_pb, f, 6, "=" + feed_formula(refs["ration"], f"C{f}"), nf="#,##0.00")
            put(ws_pb, f, 7, f"=C{f}*{refs['ops']['rate_cell']}", nf="#,##0.00")
            put(ws_pb, f, 8, f"=SUM(E{f}:G{f})", nf="#,##0.00"); put(ws_pb, f, 9, f"=D{f}-H{f}", nf="#,##0.00")
            f += 1
        t = f
        put(ws_pb, t, 1, f"=COUNTA(A{r + 2}:A{t - 1})", bold=True, fill=YELLOW)
        for c_ in range(2, 9):
            if c_ != 3:
                put(ws_pb, t, c_, f"=SUM({L(c_)}{r + 2}:{L(c_)}{t - 1})", bold=True, fill=YELLOW, nf="#,##0.00")
        put(ws_pb, t, 9, f"=D{t}-H{t}", bold=True, fill=YELLOW, nf="#,##0.00")
        summary.append((sb.batch_number, t))
        r = t + 3
    # culls & mortalities: cost with no revenue (absent from the original, which hid these losses)
    cull = [a for rows in by_batch.values() for a in rows if a.exit_type in ("culled", "died")]
    cull_t = None
    if cull:
        put(ws_pb, r, 4, "Culls and mortalities (cost, no revenue)", bold=True, border=False)
        for i, h in enumerate(["Animal No.", "Exit", "Days on Farm", "Selling price", "Buying price", "Feeding cost", "Operations cost", "Total expenses", "Profit"], 1):
            hdr(ws_pb, r + 1, i, h)
        f = r + 2
        for a in sorted(cull, key=lambda x: _nat(x.tag)):
            p = prof.get(a.animal_id, {})
            put(ws_pb, f, 1, a.tag); put(ws_pb, f, 2, a.exit_type); put(ws_pb, f, 3, p.get("days_on_farm", 0)); put(ws_pb, f, 4, 0)
            put(ws_pb, f, 5, f"={q(' PROFITABILITY ')}!H{buy_row.get(a.intake_batch, 3)}", nf="#,##0.00")
            put(ws_pb, f, 6, "=" + feed_formula(refs["ration"], f"C{f}"), nf="#,##0.00")
            put(ws_pb, f, 7, f"=C{f}*{refs['ops']['rate_cell']}", nf="#,##0.00")
            put(ws_pb, f, 8, f"=SUM(E{f}:G{f})", nf="#,##0.00"); put(ws_pb, f, 9, f"=D{f}-H{f}", nf="#,##0.00")
            f += 1
        cull_t = f
        put(ws_pb, f, 1, f"=COUNTA(A{r + 2}:A{f - 1})", bold=True, fill=YELLOW)
        for c_ in (4, 5, 6, 7, 8):
            put(ws_pb, f, c_, f"=SUM({L(c_)}{r + 2}:{L(c_)}{f - 1})", bold=True, fill=YELLOW, nf="#,##0.00")
        put(ws_pb, f, 9, f"=D{f}-H{f}", bold=True, fill=YELLOW, nf="#,##0.00")
    pbs = q("Profitability per Bull")
    rr = s_top + 2
    for n, t in summary:
        put(ws, rr, 1, f"Sales Batch {n}")
        put(ws, rr, 2, f"={pbs}!A{t}"); put(ws, rr, 3, f"={pbs}!B{t}")
        for c_, col in zip(range(4, 9), "DEFGH"):
            put(ws, rr, c_, f"={pbs}!{col}{t}", nf="#,##0.00")
        put(ws, rr, 9, f"=D{rr}-H{rr}", nf="#,##0.00")
        rr += 1
    if cull_t:
        put(ws, rr, 1, "Culls & mortalities"); put(ws, rr, 2, f"={pbs}!A{cull_t}"); put(ws, rr, 3, 0)
        for c_, col in zip(range(4, 9), "DEFGH"):
            put(ws, rr, c_, f"={pbs}!{col}{cull_t}", nf="#,##0.00")
        put(ws, rr, 9, f"=D{rr}-H{rr}", nf="#,##0.00")
        rr += 1
    put(ws, rr, 1, "TOTAL", bold=True, fill=YELLOW)
    first = s_top + 2
    for c_ in range(2, 9):
        put(ws, rr, c_, f"=SUM({L(c_)}{first}:{L(c_)}{rr - 1})", bold=True, fill=YELLOW, nf="#,##0.00" if c_ > 3 else None)
    put(ws, rr, 9, f"=D{rr}-H{rr}", bold=True, fill=YELLOW, nf="#,##0.00")
    ws.conditional_formatting.add(f"I{first}:I{rr}", FormulaRule(formula=[f"I{first}<0"], font=Font(name="Arial", color="FFC00000", bold=True)))
    ws_pb.conditional_formatting.add(f"I1:I{ws_pb.max_row}", FormulaRule(formula=["AND(ISNUMBER(I1),I1<0)"], font=Font(name="Arial", color="FFC00000")))
    for sh in (ws, ws_pb):
        sh.column_dimensions["A"].width = 18
        for c_ in range(2, 10):
            sh.column_dimensions[L(c_)].width = 16


# --------------------------------------------------------------- summary sheet
def build_summary(wb, b: Bundle, batch_refs: dict[int, dict]) -> None:
    ws = wb.create_sheet("SUMMARY ADG")
    put(ws, 4, 3, "SUMMARY OF AVERAGE DAILY GAIN ", bold=True, border=False)
    asof = max((r["as_of"] for r in batch_refs.values() if r.get("as_of")), default=b.as_of)
    for i, h in enumerate(["Date Received", "No of Bulls", "Days on Farm", f"ADG (Kgs) {asof.day}/{asof.month}/{asof.year}", "% WG", "Cost per Bull (landed)"], 3):
        hdr(ws, 5, i, h)
    off = int(b.s["day_offset"])
    r = 6
    for bt in b.batches:
        ref = batch_refs.get(bt.batch_number)
        put(ws, r, 2, f"Batch {bt.batch_number}" + (f" ({bt.source})" if bt.source else ""))
        put(ws, r, 3, bt.delivered_on, nf=DATE_FMT)
        if ref and ref.get("overall_col"):
            sh = q(ref["sheet"])
            put(ws, r, 4, f"=COUNT({sh}!{ref['overall_col']}{ref['first_row']}:{ref['overall_col']}{ref['last_row']})")
            put(ws, r, 5, f"={dexpr(asof)}-C{r}-{off}")
            put(ws, r, 6, f"={sh}!{ref['adg_cell']}", nf="0.00"); put(ws, r, 7, f"={sh}!{ref['pct_cell']}", nf="0.0%")
        else:
            for c_ in (4, 5, 6, 7):
                put(ws, r, c_, None)
        put(ws, r, 8, bt.landed_cost_per_head, nf="#,##0.00")
        r += 1
    put(ws, r, 4, f"=SUM(D6:D{r - 1})", bold=True, fill=YELLOW)
    for c_, w in zip("BCDEFGH", (18, 14, 12, 13, 18, 10, 20)):
        ws.column_dimensions[c_].width = w


# ------------------------------------------------------------ analytics sheets
def _table(ws, r0: int, headers: list[str], rows: list[list], fmts: dict[int, str] | None = None) -> int:
    for i, h in enumerate(headers, 1):
        hdr(ws, r0, i, h)
    for j, row in enumerate(rows, 1):
        for i, v in enumerate(row, 1):
            put(ws, r0 + j, i, v, nf=(fmts or {}).get(i))
    for i in range(1, len(headers) + 1):
        ws.column_dimensions[L(i)].width = max(12, min(30, len(headers[i - 1]) + 2))
    return r0 + len(rows)


def build_analytics(wb, b: Bundle) -> None:
    P, A = b.profit, b.perf
    ws = wb.create_sheet("ANALYTICS - Batches")
    put(ws, 1, 1, f"INTAKE BATCH SCORECARD (values as at {b.as_of.strftime('%d/%m/%Y')} - realised sales only; open animals shown as projection)", bold=True, border=False)
    rows = [[f"BATCH {x['intake_batch']}", x["head_total"], x["head_sold"], x["head_culled"], x["head_died"], x["head_active"], x["loss_rate_pct"] / 100,
             x["avg_weight_kg"], x["avg_adg"], x["avg_days_on_farm"], x["cost_per_kg_gain"], x["revenue"], x["total_cost"], x["profit"], x["profit_per_head"],
             x["unrealised_projected_profit"]] for x in P["by_intake_batch"]]
    _table(ws, 3, ["Batch", "Head", "Sold", "Culled", "Died", "Active", "Loss rate", "Avg exit wt (kg)", "Avg ADG", "Avg days", "Cost per kg gain",
                   "Revenue", "Total cost", "Realised profit", "Profit / head", "Unrealised projected profit"], rows,
           {7: "0.0%", 9: "0.00", 11: "#,##0.00", 12: "#,##0", 13: "#,##0", 14: "#,##0", 15: "#,##0", 16: "#,##0"})
    ws2 = wb.create_sheet("ANALYTICS - Sales")
    put(ws2, 1, 1, "SALE BATCH SCORECARD", bold=True, border=False)
    prices = {sb.batch_number: sb.price_per_kg for sb in b.sale_batches}
    rows = [[f"Sales Batch {x['sale_batch']}", x["head"], x["avg_weight_kg"], float(prices[x["sale_batch"]]) if prices.get(x["sale_batch"]) else x["effective_price_per_kg"],
             x["revenue"], x["total_cost"], x["profit"], x["profit_per_head"], x["avg_days_on_farm"], x["avg_adg"], x["avg_presale_shrink_kg"],
             round(x["total_cost"] / x["total_weight_kg"], 2) if x["total_weight_kg"] else None] for x in P["by_sale_batch"]]
    _table(ws2, 3, ["Sale batch", "Head", "Avg sale wt (kg)", "Price per kg", "Revenue", "Total cost", "Profit", "Profit / head", "Avg days", "Avg ADG",
                    "Avg pre-sale shrink (kg)", "Break-even price per kg"], rows, {4: "0.00", 5: "#,##0", 6: "#,##0", 7: "#,##0", 8: "#,##0", 10: "0.00", 12: "0.00"})
    ws3 = wb.create_sheet("ANALYTICS - Action list")
    put(ws3, 1, 1, f"ACTIVE ANIMALS - what needs attention (as at {b.as_of.strftime('%d/%m/%Y')}; market-ready = {b.s['market_ready_kg']} kg)", bold=True, border=False)
    order = {"poor": 0, "below_target": 1, "on_target": 2, "excellent": 3, None: 4}
    act = [x for x in A["animals"] if x["status"] == "active"]
    act.sort(key=lambda x: (not x["market_ready"], order.get(x["adg_band"], 4), x["tag"]))
    rows = [[x["tag"], f"BATCH {x['intake_batch']}", x["latest_weight"], x["latest_weigh_date"], x["days_on_farm"], x["adg"], (x["adg_band"] or "").replace("_", " "),
             "READY" if x["market_ready"] else "", x["est_days_to_market"], ", ".join(x["flags"])] for x in act]
    end = _table(ws3, 3, ["Animal No.", "Batch", "Latest weight (kg)", "Last weighed", "Days on farm", "ADG", "ADG band", "Market-ready", "Est. days to market", "Flags"], rows, {6: "0.00"})
    if rows:
        ws3.conditional_formatting.add(f"G4:G{end}", FormulaRule(formula=['G4="poor"'], fill=RED))
        ws3.conditional_formatting.add(f"G4:G{end}", FormulaRule(formula=['G4="on target"'], fill=GREEN))
        ws3.conditional_formatting.add(f"G4:G{end}", FormulaRule(formula=['G4="excellent"'], fill=CYAN))
        ws3.conditional_formatting.add(f"H4:H{end}", FormulaRule(formula=['H4="READY"'], fill=GREEN))


def build_notes(wb, b: Bundle, corrections: list[str]) -> None:
    ws = wb.create_sheet("DATA NOTES")
    put(ws, 1, 1, "HOW THIS WORKBOOK DIFFERS FROM THE MANUAL VERSION", bold=True, border=False)
    notes = corrections + ["",
        f"Thresholds used: market-ready >= {b.s['market_ready_kg']} kg; ADG red <= {b.s['adg_poor_max']}, green {b.s['adg_target_min']}-{b.s['adg_target_max']}, cyan > {b.s['adg_target_max']}; interval gain cyan > {b.s['interval_gain_good']} kg.",
        f"Days on farm = date difference - {b.s['day_offset']} (the worksheet's convention).",
        "Sheets after 'Profitability per Bull'/' PROFITABILITY ' (ANALYTICS - ..., DATA NOTES) are additions; every original sheet keeps its name, order and columns.",
        "This file can be imported back into the system: the importer reads BATCH n, Sales weight Batch n, the register and the costing sheets."]
    for i, n in enumerate(notes, 3):
        put(ws, i, 1, n, border=False)
    ws.column_dimensions["A"].width = 160


CORRECTIONS = [
    "1. Dates are real dates shown as dd/mm/yyyy (the manual file used m/d/yy, which swapped day and month, e.g. Batch 1 delivery read as 12 Nov instead of 11 Dec).",
    "2. Day counts, bull counts and the as-at date are calculated from the data, not typed in (no more '/228' or '/221/14').",
    "3. Total-row % WG divides by the entry weight of the SAME animals that have a gain (the manual total divided by every bull ever in the batch, showing ~15-42% instead of ~49-59%).",
    "4. Interval-gain formulas always reference their own row (the manual BATCH 1 had D12 pointing at row 4).",
    "5. Profitability totals count every sold bull (the manual count formula stopped at sales batch 4: 104 instead of 109) and a bull listed as sold twice is counted once.",
    "6. Operations cost is allocated per bull-day (pool / actual bull-days), so a 156-day bull no longer pays the same as a 237-day bull.",
    "7. Feeding cost is a live formula from the ration plan and costing sheets, not a pasted number.",
    "8. Culled and dead animals appear as cost-with-no-revenue lines instead of disappearing from profit.",
    "9. Headers describe the data they sit over (stale labels such as '% WG as at 11/5/2026' over 28/7 data are generated from the real date).",
    "10. Added: pre-sale shrink columns on every sales sheet, price/amount columns, analytics sheets, % WG on the summary, freeze panes.",
]


# --------------------------------------------------------------------- driver
async def export_worksheet(db: AsyncSession, org_id: uuid.UUID, as_of: date | None = None) -> tuple[bytes, dict]:
    b = await gather(db, org_id, as_of)
    wb = Workbook()
    wb.remove(wb.active)
    by_batch: dict[int, list[metrics.AnimalInput]] = defaultdict(list)
    by_sale: dict[int, list[metrics.AnimalInput]] = defaultdict(list)
    for inp, _ in b.items:
        by_batch[inp.intake_batch].append(inp)
        if inp.exit_type == "sold" and inp.sale_batch:
            by_sale[inp.sale_batch].append(inp)
    for bt in b.batches:
        by_batch.setdefault(bt.batch_number, [])
    # sheets are created in the original order; builders that need cross-refs run first, then sheets are re-ordered
    batch_refs = {bt.batch_number: build_batch_sheet(wb, b, bt, by_batch[bt.batch_number]) for bt in b.batches}
    sales_refs = {sb.batch_number: build_sales_sheet(wb, b, sb, by_sale[sb.batch_number], {}) for sb in b.sale_batches if by_sale.get(sb.batch_number)}
    conc = build_concentrate(wb, b)
    build_land(wb, b)
    silage = build_silage(wb, b)
    hay = build_hay(wb, b)
    ration = build_ration(wb, b, {"hay": hay["cost_cell"], "silage": silage["cost_cell"], "concentrate": conc["cost_cell"]})
    bull_days = int(b.profit["assumptions"]["total_bull_days"])
    ops = build_ops(wb, b, bull_days)
    build_profitability(wb, b, by_batch, by_sale, {"ration": ration, "ops": ops})
    build_register(wb, b, {n: v for n, v in by_batch.items() if v})
    build_summary(wb, b, batch_refs)
    build_analytics(wb, b)
    build_notes(wb, b, CORRECTIONS)
    original = ["SUMMARY ADG", "Bull Inventory register"] + [sales_sheet_name(n) for n in sorted(sales_refs)] + \
               [intake_sheet_name(bt.batch_number) for bt in b.batches] + ["CONCENTRATE", "LAND USE", "SILAGE -BLENDED", "HAY",
               "TOTAL FEEDING COSTS", "OPERATIONS COST", "Profitability per Bull", " PROFITABILITY "]
    extras = ["ANALYTICS - Batches", "ANALYTICS - Sales", "ANALYTICS - Action list", "DATA NOTES"]
    wb._sheets = [wb[n] for n in original + extras if n in wb.sheetnames]
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue(), {"as_of": b.as_of.isoformat(), "sheets": wb.sheetnames, "animals": len(b.items)}
