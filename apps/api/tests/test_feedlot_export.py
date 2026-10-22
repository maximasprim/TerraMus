"""Export of the Sahara Oasis worksheet: exact layout, correct live formulas, round-trip import."""
import io
import pathlib
import shutil
import subprocess
import tempfile

import openpyxl
import pytest
from sqlalchemy import text

pytestmark = pytest.mark.asyncio
FIXTURE = pathlib.Path(__file__).parent / "fixtures" / "sahara_oasis_bulls_wg_sample.xlsx"
ORIGINAL_SHEETS = ["SUMMARY ADG", "Bull Inventory register", " Sales weight Batch 1", " Sales weight Batch 2", "Sales weight Batch 3",
                   "Sales weight Batch4", "Sales Weight Batch 5", "Sales  weight Batch 6 ", "BATCH 1", "BATCH 2", "BATCH 3", "BATCH 4.",
                   "BATCH 5", "CONCENTRATE", "LAND USE", "SILAGE -BLENDED", "HAY", "TOTAL FEEDING COSTS", "OPERATIONS COST",
                   "Profitability per Bull", " PROFITABILITY "]


async def _org(client, engine, email):
    reg = await client.post("/api/v1/auth/register", json={
        "organization_name": "Export Farm", "full_name": "Owner One", "email": email, "password": "supersecret123"})
    h = {"Authorization": f"Bearer {reg.json()['access_token']}"}
    await client.post("/api/v1/farms", json={"name": "Main"}, headers=h)
    async with engine.begin() as c:
        await c.execute(text("UPDATE subscription_tiers SET max_animals = NULL"))
    return h


async def _import(client, h, content, name="w.xlsx"):
    r = await client.post("/api/v1/feedlot/worksheet/import?dry_run=false", headers=h,
                          files={"file": (name, io.BytesIO(content), "application/octet-stream")})
    assert r.status_code == 200, r.text
    return r.json()


async def _imported_and_exported(client, engine, email):
    h = await _org(client, engine, email)
    rep = await _import(client, h, FIXTURE.read_bytes())
    assert rep["ok"]
    e = await client.get("/api/v1/feedlot/worksheet/export", headers=h)
    assert e.status_code == 200 and "spreadsheetml" in e.headers["content-type"]
    return h, rep, e.content


async def test_export_keeps_original_sheet_names_order_and_adds_extras_after(client, engine):
    _, _, content = await _imported_and_exported(client, engine, "ex1@example.com")
    names = openpyxl.load_workbook(io.BytesIO(content)).sheetnames
    assert names[:21] == ORIGINAL_SHEETS
    assert names[21:] == ["ANALYTICS - Batches", "ANALYTICS - Sales", "ANALYTICS - Action list", "DATA NOTES"]


async def test_export_layout_matches_worksheet_conventions(client, engine):
    _, _, content = await _imported_and_exported(client, engine, "ex2@example.com")
    wb = openpyxl.load_workbook(io.BytesIO(content))
    b1 = wb["BATCH 1"]
    assert b1["A1"].value == "BATCH 1. Delivered on 11/12/2025." and b1["A2"].value == "ANIMAL NO."
    assert b1["B2"].value.date().isoformat() == "2025-12-11" and b1["B2"].number_format == "dd/mm/yyyy"
    assert "Day Gain/ (Loss)" in b1["D2"].value
    assert b1["A3"].value == "A01" and b1["A33"].value == "TOTAL"
    assert b1["A2"].fill.fgColor.rgb == "00FFFF00" or b1["A2"].fill.fgColor.rgb == "FFFFFF00"
    assert wb["BATCH 4."]["A1"].value.startswith("BATCH 4. ISIOLO- 30 Steers.")
    sales = wb[" Sales weight Batch 1"]
    assert [c.value for c in sales[1]][:8] == ["Animal No.", "Date received", "Entry weight(Kg)", "Date sold", "Sales Weight(Kg)", "TWG", "Days on Farm", "ADG"]
    reg = wb["Bull Inventory register"]
    assert reg["B1"].value == "SOLD BULLS" and reg["A2"].value == "Batch 1" and str(reg["A3"].value).endswith("...1")


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
async def test_exported_formulas_calculate_without_errors_and_match_worksheet(client, engine):
    _, _, content = await _imported_and_exported(client, engine, "ex3@example.com")
    with tempfile.TemporaryDirectory() as td:
        src = pathlib.Path(td) / "x.xlsx"; src.write_bytes(content)
        out = pathlib.Path(td) / "out"
        subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(out), str(src)],
                       check=True, capture_output=True, timeout=120)
        wb = openpyxl.load_workbook(out / "x.xlsx", data_only=True)
    errors = [(ws.title, c.coordinate, c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row
              if isinstance(c.value, str) and (c.value.startswith("#") or c.value.startswith("Err:"))]
    assert errors == []
    s = wb["SUMMARY ADG"]
    assert [s.cell(r, 5).value for r in range(6, 11)] == [228, 221, 215, 199, 192]          # worksheet's days on farm
    assert [s.cell(r, 4).value for r in range(6, 11)] == [11, 14, 11, 24, 5]                # worksheet's bulls weighed 28/7
    assert [round(s.cell(r, 7).value, 3) for r in range(6, 11)] == [0.540, 0.495, 0.585, 0.549, 0.491]   # corrected % WG
    p = wb[" PROFITABILITY "]
    labels = [p.cell(r, 1).value for r in range(1, 30)]
    sb1 = labels.index("Sales Batch 1") + 1
    assert p.cell(sb1, 4).value == 3623070 and p.cell(sb1, 5).value == pytest.approx(1690652, abs=1)   # same as the worksheet
    total_row = labels.index("TOTAL") + 1
    assert p.cell(total_row, 2).value == 109 and "Culls & mortalities" in labels
    assert p.cell(total_row, 9).value == pytest.approx(p.cell(total_row, 4).value - p.cell(total_row, 8).value, abs=1)


async def test_exported_file_imports_back_to_the_same_data(client, engine):
    h1, rep1, content = await _imported_and_exported(client, engine, "ex4@example.com")
    h2 = await _org(client, engine, "ex5@example.com")
    dry = await client.post("/api/v1/feedlot/worksheet/import?dry_run=true", headers=h2,
                            files={"file": ("e.xlsx", io.BytesIO(content), "application/octet-stream")})
    assert dry.json()["created"].get("animals") == 131
    rep2 = await _import(client, h2, content, "e.xlsx")
    assert rep2["ok"], rep2["errors"]
    for k in ("intake_batches", "sale_batches", "animals", "weight_records", "sales", "culls", "costing_documents"):
        assert rep2["created"].get(k) == rep1["created"].get(k), (k, rep1["created"].get(k), rep2["created"].get(k))
    a = (await client.get("/api/v1/feedlot/profitability?as_of=2026-09-30", headers=h1)).json()["totals"]
    b = (await client.get("/api/v1/feedlot/profitability?as_of=2026-09-30", headers=h2)).json()["totals"]
    assert a["revenue"] == b["revenue"] and a["profit"] == pytest.approx(b["profit"], abs=1)
    # exporting the re-imported data gives the same worksheet shape
    e2 = await client.get("/api/v1/feedlot/worksheet/export", headers=h2)
    assert openpyxl.load_workbook(io.BytesIO(e2.content)).sheetnames == openpyxl.load_workbook(io.BytesIO(content)).sheetnames


async def test_export_of_empty_org_is_a_valid_workbook_and_requires_auth(client, engine):
    reg = await client.post("/api/v1/auth/register", json={
        "organization_name": "Empty Farm", "full_name": "Owner One", "email": "ex6@example.com", "password": "supersecret123"})
    h = {"Authorization": f"Bearer {reg.json()['access_token']}"}
    e = await client.get("/api/v1/feedlot/worksheet/export", headers=h)
    assert e.status_code == 200 and openpyxl.load_workbook(io.BytesIO(e.content)).sheetnames[0] == "SUMMARY ADG"
    assert (await client.get("/api/v1/feedlot/worksheet/export")).status_code in (401, 403)
    bad = await client.post("/api/v1/feedlot/worksheet/import", headers=h, files={"file": ("x.txt", io.BytesIO(b"hi"), "text/plain")})
    assert bad.status_code == 422
    junk = await client.post("/api/v1/feedlot/worksheet/import", headers=h, files={"file": ("x.xlsx", io.BytesIO(b"not excel"), "application/octet-stream")})
    assert junk.status_code == 422
